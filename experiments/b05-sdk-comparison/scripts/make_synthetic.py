#!/usr/bin/env python3
"""Generate entirely fictional DOCX controls. Never transform NITE official models."""
import argparse
from datetime import datetime,timezone
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
import struct,zlib
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def complex_field(paragraph, instruction, cache):
    for tag,value in (('fldChar','begin'),('instrText',instruction),('fldChar','separate'),('t',cache),('fldChar','end')):
        run=OxmlElement('w:r');element=OxmlElement('w:'+tag)
        if tag=='fldChar':element.set(qn('w:fldCharType'),value)
        else:element.text=value
        run.append(element);paragraph._p.append(run)

def synthetic_logo():
    def chunk(kind,data):return struct.pack('!I',len(data))+kind+data+struct.pack('!I',zlib.crc32(kind+data))
    scan=b''.join(b'\0'+bytes((30,80,130))*80 for _ in range(40))
    return b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('!2I5B',80,40,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(scan))+chunk(b'IEND',b'')

def make(path:Path,*,extended=False):
    doc=Document()
    doc.core_properties.modified=datetime(2026,9,23,23,56,0,tzinfo=timezone.utc)
    doc.add_heading('NITE — fixture sintética de compatibilidade',level=1)
    for i in range(4):
        p=doc.add_paragraph(f'Parágrafo fictício {i+1} com espaçamento controlado.')
        p.paragraph_format.space_after=Pt(8)
    doc.add_paragraph('A estrutura a seguir é uma tabela anônima.')
    t=doc.add_table(rows=3,cols=3)
    for row in range(3):
        for col in range(3):t.cell(row,col).text=f'F{row}-{col}'
    s=doc.sections[0]
    header=s.header.paragraphs[0]
    header.add_run('Última revisão: ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'SAVEDATE  \\* MERGEFORMAT')
    r=OxmlElement('w:r');txt=OxmlElement('w:t');txt.text='23/09/2026 23:56:00';r.append(txt);fld.append(r);header._p.append(fld)
    p=s.footer.paragraphs[0];p.add_run('Página ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE \\* Arabic');r2=OxmlElement('w:r');t2=OxmlElement('w:t');t2.text='1';r2.append(t2);f.append(r2);p._p.append(f)
    if extended:
        header.add_run(' · Campo complexo: ');complex_field(header,' SAVEDATE \\* MERGEFORMAT ','23/09/2026 23:56:00')
        header.add_run().add_picture(BytesIO(synthetic_logo()),width=Inches(0.6))
        p.add_run(' de ');complex_field(p,' NUMPAGES ','2')
        doc.add_page_break();doc.add_paragraph('Segunda página sintética com quebra explícita e campos de cabeçalho e rodapé.')
    path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(path)
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--out',type=Path,required=True);parser.add_argument('--extended',action='store_true');arg=parser.parse_args()
    if '.private' not in arg.out.resolve().parts:raise SystemExit('Generate only into ignored .private')
    if arg.out.exists():raise SystemExit('Existing fixture: refusing to overwrite')
    make(arg.out,extended=arg.extended);print('Synthetic fixture created locally')
