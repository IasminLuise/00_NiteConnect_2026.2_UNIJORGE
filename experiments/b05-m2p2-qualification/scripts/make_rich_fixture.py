#!/usr/bin/env python3
"""Generate a fictional multi-page DOCX with table, image, fields, header/footer and footnote.

Synthetic testing only. It is never a sanitized copy of an official document.
"""
from __future__ import annotations
import argparse
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import struct,zlib
from lxml import etree as ET
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
W='{'+NS+'}'


def logo():
    def png_chunk(kind,raw):return struct.pack('>I',len(raw))+kind+raw+struct.pack('>I',zlib.crc32(kind+raw)&0xffffffff)
    rows=b''.join(b'\0'+bytes([40,110,170])*90 for _ in range(40))
    return b'\x89PNG\r\n\x1a\n'+png_chunk(b'IHDR',struct.pack('>IIBBBBB',90,40,8,2,0,0,0))+png_chunk(b'IDAT',zlib.compress(rows))+png_chunk(b'IEND',b'')


def field(p,name,complex_=False):
    if complex_:
        def component(tag,attributes=None,text=None):
            r=OxmlElement('w:r');node=OxmlElement('w:'+tag)
            for attr,value in (attributes or {}).items():node.set(qn('w:'+attr),value)
            if text is not None:node.text=text
            r.append(node);p._p.append(r)
        component('fldChar',{'fldCharType':'begin'});component('instrText',text=' '+name+' \\* MERGEFORMAT ')
        component('fldChar',{'fldCharType':'separate'});component('t',text='1');component('fldChar',{'fldCharType':'end'})
    else:
        e=OxmlElement('w:fldSimple');e.set(qn('w:instr'),name)
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);e.append(r);p._p.append(e)


def inject_footnote(docx:Path):
    with ZipFile(docx) as z: parts={name:z.read(name) for name in z.namelist()}
    parser=ET.XMLParser(resolve_entities=False,no_network=True)
    doc=ET.fromstring(parts['word/document.xml'],parser)
    body=doc.find(W+'body')
    paragraphs=list(body.iter(W+'p'))
    if len(paragraphs)<2:raise ValueError('missing paragraph')
    run=ET.SubElement(paragraphs[1],W+'r');ET.SubElement(run,W+'footnoteReference',{W+'id':'1'})
    parts['word/document.xml']=ET.tostring(doc,encoding='UTF-8',xml_declaration=True,standalone=True)
    notes=ET.Element(W+'footnotes',nsmap={'w':NS})
    for id_,ty,marker in (('-1','separator','separator'),('0','continuationSeparator','continuationSeparator')):
        node=ET.SubElement(notes,W+'footnote',{W+'type':ty,W+'id':id_})
        p=ET.SubElement(node,W+'p');r=ET.SubElement(p,W+'r');ET.SubElement(r,W+marker)
    user=ET.SubElement(notes,W+'footnote',{W+'id':'1'})
    p=ET.SubElement(user,W+'p');run=ET.SubElement(p,W+'r');ET.SubElement(run,W+'t').text='Nota inteiramente fictícia da fixture B05.'
    parts['word/footnotes.xml']=ET.tostring(notes,encoding='UTF-8',xml_declaration=True,standalone=True)
    rels=ET.fromstring(parts['word/_rels/document.xml.rels'],parser)
    ids={node.get('Id') for node in rels}
    n=5000
    while 'rId'+str(n) in ids:n+=1
    ET.SubElement(rels,'{'+REL+'}Relationship',{'Id':'rId'+str(n),'Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes','Target':'footnotes.xml'})
    parts['word/_rels/document.xml.rels']=ET.tostring(rels,encoding='UTF-8',xml_declaration=True,standalone=True)
    types=ET.fromstring(parts['[Content_Types].xml'],parser)
    ET.SubElement(types,'{'+CT+'}Override',{'PartName':'/word/footnotes.xml','ContentType':'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'})
    parts['[Content_Types].xml']=ET.tostring(types,encoding='UTF-8',xml_declaration=True,standalone=True)
    tmp=docx.with_suffix('.tmp')
    with ZipFile(tmp,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    tmp.replace(docx)


def make(path:Path):
    if '.private' not in path.resolve().parts or path.exists():raise ValueError('private new path only')
    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
    style=doc.styles['Normal'];style.font.name='Cambria';style.font.size=Pt(11)
    doc.add_heading('Relatório inteiramente sintético de fidelidade',1)
    doc.add_paragraph('Contexto fictício acompanhado de uma nota de rodapé.')
    doc.add_paragraph('MARCADOR_EDITAVEL_M2P2 — versão inicial controlada.')
    table=doc.add_table(rows=3,cols=3)
    for row in range(3):
        for col in range(3):table.cell(row,col).text=f'Item fictício {row+1}.{col+1}'
    doc.add_paragraph('Imagem sintética de teste:').add_run().add_picture(BytesIO(logo()),width=Inches(1.2))
    header=sec.header.paragraphs[0];header.add_run('Última gravação: ');field(header,'SAVEDATE',complex_=True)
    footer=sec.footer.paragraphs[0];footer.add_run('Página ');field(footer,'PAGE');footer.add_run(' de ');field(footer,'NUMPAGES',complex_=True)
    doc.add_page_break()
    doc.add_heading('Segunda página fictícia',2)
    for i in range(10):
        p=doc.add_paragraph(f'Parágrafo de validação {i+1}: conteúdo fictício sem qualquer informação institucional.')
        p.paragraph_format.space_after=Pt(8)
    path.parent.mkdir(parents=True,exist_ok=True);doc.save(path);inject_footnote(path)
    return path


if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--out',required=True,type=Path);opts=a.parse_args()
    make(opts.out);print('Synthetic fixture created only in private storage')
