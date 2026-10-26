import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import fitz
from lxml import etree as ET
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from semantic import analyze_semantics,style_map,compare_styles,relationships,content_types,ref_targets
from visual import compare_pdfs
from policy import decide
from make_rich_fixture import make
from docx import Document

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class Qualification(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp=tempfile.TemporaryDirectory()
        cls.root=Path(cls._tmp.name)/'.private';cls.root.mkdir()
        cls.source=cls.root/'rich.docx';make(cls.source)
    @classmethod
    def tearDownClass(cls):cls._tmp.cleanup()
    def setUp(self):
        self.candidate=self.root/(self.id().rsplit('.',1)[-1]+'.docx')
        self.expected={'paragraphIndex':2,'oldFragment':'MARCADOR_EDITAVEL_M2P2','newFragment':'MARCADOR_EDITADO_M2P2'}
        self.edit()
    def repack(self,mutator):
        with ZipFile(self.source) as z:parts={n:z.read(n) for n in z.namelist()}
        mutator(parts)
        with ZipFile(self.candidate,'w',ZIP_DEFLATED) as z:
            for n,d in parts.items():z.writestr(n,d)
    def edit(self):
        def edit_text(parts):
            old=b'MARCADOR_EDITAVEL_M2P2';new=b'MARCADOR_EDITADO_M2P2'
            self.assertIn(old,parts['word/document.xml']);parts['word/document.xml']=parts['word/document.xml'].replace(old,new,1)
        self.repack(edit_text)
    def analyze(self):return analyze_semantics(self.source,self.candidate,self.expected)
    def test_fixture_has_multiple_pages_header_footer_media_table_footnote_fields(self):
        d=Document(self.source)
        self.assertEqual(len(d.tables),1)
        self.assertIn('MARCADOR_EDITAVEL_M2P2',' '.join(p.text for p in d.paragraphs))
        with ZipFile(self.source) as z:
            names=set(z.namelist())
            self.assertIn('word/footnotes.xml',names)
            self.assertTrue(any(n.startswith('word/media/') for n in names))
            self.assertIn('word/header1.xml',names);self.assertIn('word/footer1.xml',names)
            self.assertIn(b'SAVEDATE',z.read('word/header1.xml'))
            self.assertIn(b'NUMPAGES',z.read('word/footer1.xml'))
    def test_pure_text_structural_prequalification(self):
        s=self.analyze()
        self.assertEqual(s['classification'],'EXACT_TEXT_ONLY')
        self.assertTrue(s['expectedEditExact'])
        self.assertEqual(s['styles']['state'],'EQUIVALENT_BY_EXACT_STYLE_DEFINITION')
        self.assertEqual(s['fields']['state'],'EQUAL')
        self.assertEqual(s['notes']['status'],'UNCHANGED')
        self.assertEqual(s['risks'],[])
        self.assertFalse(s['mayPromote'])
    def test_wrong_text_expected_fails(self):
        self.expected['newFragment']='MUDANCA_INCORRETA'
        self.assertFalse(self.analyze()['expectedEditExact'])
    def test_new_empty_notes_are_not_authorized(self):
        def modify(d):
            d['word/endnotes.xml']=b'<w:endnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        self.repack(modify)
        s=self.analyze()
        self.assertNotEqual(s['classification'],'EXACT_TEXT_ONLY')
        self.assertEqual(s['notes']['status'],'NEW_EMPTY_NOTES_NOT_APPROVED')
    def test_note_user_content_changed_holds(self):
        def modify(d):d['word/footnotes.xml']=d['word/footnotes.xml'].replace(b'Nota inteiramente',b'Nota totalmente')
        self.repack(modify)
        self.assertEqual(self.analyze()['notes']['status'],'REQUIRES_REVIEW')
    def test_visual_content_in_separator_is_not_empty(self):
        for tag in ('sym','pict','unknownVisual'):
            with self.subTest(tag=tag):
                def modify(d):
                    root=ET.Element(W+'endnotes',nsmap={'w':W[1:-1]})
                    note=ET.SubElement(root,W+'endnote',{W+'type':'separator',W+'id':'-1'})
                    p=ET.SubElement(note,W+'p');r=ET.SubElement(p,W+'r')
                    ET.SubElement(r,W+tag)
                    d['word/endnotes.xml']=ET.tostring(root)
                self.repack(modify)
                s=self.analyze()
                self.assertFalse(s['notes']['candidateSeparatorOnly'])
                self.assertEqual(s['notes']['status'],'REQUIRES_REVIEW')
    def test_known_separator_formatting_is_empty_but_not_approved(self):
        def modify(d):
            root=ET.Element(W+'endnotes',nsmap={'w':W[1:-1]})
            note=ET.SubElement(root,W+'endnote',{W+'type':'separator',W+'id':'-1'})
            p=ET.SubElement(note,W+'p');props=ET.SubElement(p,W+'pPr')
            ET.SubElement(props,W+'spacing');ET.SubElement(props,W+'ind')
            r=ET.SubElement(p,W+'r');ET.SubElement(r,W+'rPr');ET.SubElement(r,W+'separator')
            d['word/endnotes.xml']=ET.tostring(root)
        self.repack(modify)
        s=self.analyze()
        self.assertTrue(s['notes']['candidateSeparatorOnly'])
        self.assertEqual(s['notes']['status'],'NEW_EMPTY_NOTES_NOT_APPROVED')
        self.assertEqual(decide(s,{},None)['decision'],'HOLD_QUARANTINE')
    def test_field_instruction_changes_holds(self):
        def modify(d):d['word/header1.xml']=d['word/header1.xml'].replace(b'SAVEDATE',b'REF')
        self.repack(modify)
        self.assertEqual(self.analyze()['fields']['state'],'DIFFERENT')
    def test_table_geometry_change_holds(self):
        def modify(d):
            x=d['word/document.xml'];self.assertIn(b'gridCol',x)
            d['word/document.xml']=x.replace(b'gridCol',b'gridCxl',1)
        self.repack(modify)
        self.assertNotEqual(self.analyze()['classification'],'EXACT_TEXT_ONLY')
    def test_media_change_holds(self):
        def modify(d):
            n=next(n for n in d if n.startswith('word/media/'));d[n]=d[n]+b'Z'
        self.repack(modify)
        self.assertTrue(self.analyze()['protectedParts']['mediaChanged'])
    def test_style_property_change(self):
        def modify(d):d['word/styles.xml']=d['word/styles.xml'].replace(b'Cambria',b'Calibri',1)
        self.repack(modify)
        self.assertEqual(self.analyze()['styles']['state'],'DIFFERENT_OR_UNPROVEN')
    def test_style_alias_preserves_style_definition(self):
        def modify(d):
            root=ET.fromstring(d['word/styles.xml'])
            style=next(e for e in root.findall(W+'style') if e.get(W+'styleId')=='Heading1')
            style.set(W+'styleId','693')
            for e in root.iter():
                if e.tag in (W+'basedOn',W+'next',W+'link') and e.get(W+'val')=='Heading1':e.set(W+'val','693')
            d['word/styles.xml']=ET.tostring(root)
            root=ET.fromstring(d['word/document.xml'])
            for e in root.iter(W+'pStyle'):
                if e.get(W+'val')=='Heading1':e.set(W+'val','693')
            d['word/document.xml']=ET.tostring(root)
        self.repack(modify)
        s=self.analyze()
        self.assertEqual(s['styles']['state'],'EQUIVALENT_BY_EXACT_STYLE_DEFINITION')
        self.assertEqual(s['classification'],'OOXML_TRANSFORMATIONS_REQUIRE_JUSTIFICATION')
    def test_added_external_relationship_fails_closed(self):
        def modify(d):
            x=ET.fromstring(d['word/_rels/document.xml.rels'])
            ET.SubElement(x,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                          {'Id':'rId500','Type':'urn:evil','Target':'https://example.org/collect','TargetMode':'External'})
            d['word/_rels/document.xml.rels']=ET.tostring(x)
        self.repack(modify)
        self.assertIn('EXTERNAL_RELATIONSHIP',self.analyze()['risks'])
    def test_dangling_relationship_reference_fails_closed(self):
        def modify(d):
            x=ET.fromstring(d['word/document.xml'])
            dr=next(e for e in x.iter() if e.tag.endswith('}blip'))
            dr.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed','rId404missing')
            d['word/document.xml']=ET.tostring(x)
        self.repack(modify)
        self.assertIn('OOXML_DANGLING_RELATIONSHIP_REFERENCE',self.analyze()['risks'])
    def test_missing_content_type_fails_closed(self):
        def modify(d):
            x=ET.fromstring(d['[Content_Types].xml'])
            for e in x:
                if e.get('PartName')=='/word/footnotes.xml':x.remove(e)
            d['[Content_Types].xml']=ET.tostring(x)
        self.repack(modify)
        self.assertIn('WRONG_PROTECTED_PART_CONTENT_TYPE',self.analyze()['risks'])
    def test_removed_thumbnail_not_inherently_allowed(self):
        def modify(d):d.pop('docProps/thumbnail.jpeg',None)
        self.repack(modify)
        s=self.analyze()
        self.assertFalse(s['mayPromote'])
        self.assertNotEqual(s['classification'],'EXACT_TEXT_ONLY')
    def test_shadow_rejects_semantic_only_even_with_good_visual(self):
        s=self.analyze()
        s['classification']='OOXML_TRANSFORMATIONS_REQUIRE_JUSTIFICATION'
        out=decide(s,{'classification':'VISUAL_MEASUREMENT_PREQUALIFIED_NOT_WORD_FIDELITY'},
                   {'decision':'SYNTHETIC_WORD_LAYOUT_PREQUALIFIED_MANUAL_VISUAL_REVIEW_REQUIRED'})
        self.assertEqual(out['decision'],'HOLD_QUARANTINE')
        self.assertFalse(out['canPromote'])
    def test_shadow_eligible_is_never_publish_approved(self):
        s=self.analyze()
        out=decide(s,{'classification':'VISUAL_MEASUREMENT_PREQUALIFIED_NOT_WORD_FIDELITY'},
                   {'decision':'SYNTHETIC_WORD_LAYOUT_PREQUALIFIED_MANUAL_VISUAL_REVIEW_REQUIRED'})
        self.assertEqual(out['decision'],'SYNTHETIC_POLICY_REVIEW_ELIGIBLE')
        self.assertFalse(out['canPromote'])
        self.assertFalse(out['officialApproved'])
    def test_shareable_cli_smoke_with_synthetic_mock_metrics_not_vendor_proof(self):
        pdf_a=self.root/'cli-original.pdf';pdf_b=self.root/'cli-edit.pdf'
        for filename,txt in ((pdf_a,'TEXTO ANTIGO'),(pdf_b,'TEXTO NOVO')):
            d=fitz.open();pg=d.new_page();pg.insert_text((75,80),txt,fontsize=11);d.save(filename);d.close()
        expected=self.root/'cli-expected.json';expected.write_text(json.dumps(self.expected),encoding='utf-8')
        summary=self.root/'cli-word-summary.json'
        summary.write_text(json.dumps({'decision':'SYNTHETIC_WORD_LAYOUT_PREQUALIFIED_MANUAL_VISUAL_REVIEW_REQUIRED'}),encoding='utf-8')
        output=self.root/'cli-shareable.json'
        cmd=[sys.executable,str(ROOT/'scripts'/'qualify.py'),'--source',str(self.source),'--candidate',str(self.candidate),
             '--expected',str(expected),'--source-pdf',str(pdf_a),'--candidate-pdf',str(pdf_b),
             '--word-summary',str(summary),'--out',str(output)]
        cp=subprocess.run(cmd,capture_output=True,text=True,timeout=15)
        self.assertEqual(cp.returncode,0,cp.stderr)
        data=json.loads(output.read_text(encoding='utf-8'))
        self.assertTrue(data['semantics']['expectedEditExact'])
        self.assertFalse(data['shadowPolicy']['canPromote'])
        self.assertNotIn('MARCADOR_EDITAVEL_M2P2',output.read_text(encoding='utf-8'))
        self.assertEqual(data['shadowPolicy']['decision'],'SYNTHETIC_POLICY_REVIEW_ELIGIBLE')
        second=subprocess.run(cmd,capture_output=True,text=True,timeout=15)
        self.assertNotEqual(second.returncode,0)
    def test_shadow_requires_word_metrics(self):
        s=self.analyze()
        out=decide(s,{'classification':'VISUAL_MEASUREMENT_PREQUALIFIED_NOT_WORD_FIDELITY'})
        self.assertEqual(out['decision'],'HOLD_QUARANTINE')
    def test_shadow_rejects_unexplained_pixels(self):
        s=self.analyze()
        out=decide(s,{'classification':'UNEXPLAINED_RESIDUAL_HOLD'},
                   {'decision':'SYNTHETIC_WORD_LAYOUT_PREQUALIFIED_MANUAL_VISUAL_REVIEW_REQUIRED'})
        self.assertEqual(out['decision'],'HOLD_QUARANTINE')
    def test_original_document_outside_private_rejected(self):
        with self.assertRaises(Exception):
            analyze_semantics(Path('/tmp/not-private.docx'),self.candidate,self.expected)

class VisualTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)/'.private';self.root.mkdir()
        self.a=self.root/'a.pdf';self.b=self.root/'b.pdf'
    def create(self,side=False,unexpected=False,second=False):
        for path,txt in ((self.a,'TEXTO ANTIGO'),(self.b,'TEXTO NOVO')):
            d=fitz.open();p=d.new_page(width=612,height=792)
            p.insert_text((75,80),txt,fontsize=11)
            p.insert_text((75,150),'Texto estavel',fontsize=11)
            if unexpected and path==self.b:p.draw_rect(fitz.Rect(200,260,204,264),color=(1,0,0),fill=(1,0,0))
            if side and path==self.b:p.insert_text((75,150),' EXTRA',fontsize=11)
            if second and path==self.b:d.new_page()
            d.save(path);d.close()
    def test_pdf_one_changed_line_no_residual(self):
        self.create()
        out=compare_pdfs(self.a,self.b)
        self.assertEqual(out['classification'],'VISUAL_MEASUREMENT_PREQUALIFIED_NOT_WORD_FIDELITY')
        self.assertEqual(out['unexplainedPixels'],0)
        self.assertEqual(out['changedTextLines'],1)
    def test_unexplained_pixels_hold(self):
        self.create(unexpected=True)
        out=compare_pdfs(self.a,self.b)
        self.assertEqual(out['classification'],'UNEXPLAINED_RESIDUAL_HOLD')
        self.assertGreater(out['unexplainedPixels'],0)
        self.assertTrue(out['pages'][0]['residualComponents'])
    def test_untouched_line_change_hold(self):
        self.create(side=True)
        out=compare_pdfs(self.a,self.b)
        self.assertNotEqual(out['classification'],'VISUAL_MEASUREMENT_PREQUALIFIED_NOT_WORD_FIDELITY')
    def test_page_count_change_holds(self):
        self.create(second=True)
        self.assertEqual(compare_pdfs(self.a,self.b)['classification'],'HOLD_QUARANTINE')
    def test_nonprivate_pdf_rejected(self):
        self.create()
        with self.assertRaises(ValueError):compare_pdfs(Path('/tmp/absent.pdf'),self.b)

if __name__=='__main__':unittest.main()
