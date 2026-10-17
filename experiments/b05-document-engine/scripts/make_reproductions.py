"""Synthetic SAVEDATE and anchor cases. Artifacts remain private pending disclosure review."""
import argparse
import copy
import io
from pathlib import Path

from docx import Document
from lxml import etree as ET
from PIL import Image

import generate_fixtures as g
from fidelity import ROOT, private_path


def generate(source, out):
    out = private_path(ROOT / '.runtime', out)
    out.mkdir(parents=True, exist_ok=True)
    date_doc = Document()
    date_doc.core_properties.author = 'Synthetic'
    date_doc.core_properties.last_modified_by = 'Synthetic'
    field = ET.SubElement(date_doc.sections[0].header.paragraphs[0]._p, '{' + g.W + '}fldSimple')
    field.set('{' + g.W + '}instr', 'SAVEDATE  \\* MERGEFORMAT')
    run = ET.SubElement(field, '{' + g.W + '}r')
    ET.SubElement(run, '{' + g.W + '}t').text = '23/09/2026 23:56:00'
    for i in range(2):
        date_doc.add_paragraph('Synthetic field page ' + str(i + 1))
        if i == 0:
            date_doc.add_page_break()
    date_doc.save(out / 'F07.docx')

    anchor_doc = Document()
    anchor_doc.core_properties.author = 'Synthetic'
    anchor_doc.core_properties.last_modified_by = 'Synthetic'
    image = io.BytesIO()
    Image.new('RGB', (80, 80), 'blue').save(image, format='PNG')
    image.seek(0)
    picture = anchor_doc.sections[0].header.paragraphs[0].add_run().add_picture(image)
    original = Document(source)
    wp = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
    a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    r = '{' + g.R + '}'
    anchors = original.sections[0].header._element.findall('.//' + wp + 'anchor')
    if not anchors:
        raise ValueError('Source does not contain the expected anchor')
    anchor = copy.deepcopy(anchors[0])
    identifier = picture._inline.find('.//' + a + 'blip').get(r + 'embed')
    for blip in anchor.findall('.//' + a + 'blip'):
        blip.set(r + 'embed', identifier)
    for prop in anchor.findall('.//' + wp + 'docPr'):
        prop.set('name', 'Synthetic anchor')
        prop.set('descr', 'Synthetic image')
    picture._inline.getparent().replace(picture._inline, anchor)
    for i in range(3):
        anchor_doc.add_paragraph('Synthetic anchor page ' + str(i + 1))
        if i < 2:
            anchor_doc.add_page_break()
    anchor_doc.save(out / 'F08.docx')
    print('Synthetic field and anchor fixtures created; do not publish without inspection')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--anchor-source', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    generate(args.anchor_source, args.out)
