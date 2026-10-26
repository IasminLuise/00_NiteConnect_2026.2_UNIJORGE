"""Synthetic, three-page field control. Never reads institutional documents."""
from io import BytesIO
from pathlib import Path
from datetime import datetime, timezone
import argparse
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw


def append_field(p, instruction, result, complex_field=False):
    if complex_field:
        for kind in ['begin', 'instruction', 'separate', 'result', 'end']:
            r = OxmlElement('w:r')
            if kind in ['begin', 'separate', 'end']:
                el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind)
            else:
                el = OxmlElement('w:instrText' if kind == 'instruction' else 'w:t')
                el.text = instruction if kind == 'instruction' else result
                el.set(qn('xml:space'), 'preserve')
            r.append(el); p._p.append(r)
    else:
        el = OxmlElement('w:fldSimple'); el.set(qn('w:instr'), instruction)
        r = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = result
        r.append(text); el.append(r); p._p.append(el)


def make(target, complex_savedate=False):
    if target.exists() or '.private' not in target.resolve().parts:
        raise ValueError('Use a new private output file')
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(11)
    doc.core_properties.modified = datetime(2020, 1, 2, 3, 4, 5, tzinfo=timezone.utc)
    section = doc.sections[0]; section.different_first_page_header_footer = True
    doc.settings.element.append(OxmlElement('w:evenAndOddHeaders'))
    logo = Image.new('RGB', (100, 60), 'white'); draw = ImageDraw.Draw(logo)
    draw.rectangle((4, 4, 96, 56), fill='#184e83'); draw.text((20, 22), 'B05', fill='white')
    stream = BytesIO(); logo.save(stream, format='PNG'); image = stream.getvalue()
    for label, header in [('PRINCIPAL', section.header), ('PRIMEIRA', section.first_page_header), ('PAR', section.even_page_header)]:
        p = header.paragraphs[0]; p.add_run().add_picture(BytesIO(image), width=Inches(.55))
        p.add_run(' ORGANIZAÇÃO FICTÍCIA — '+label)
        p = header.add_paragraph('Última revisão: ')
        append_field(p, 'SAVEDATE  \\* MERGEFORMAT', '02/01/2020 03:04:05', complex_savedate)
    for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
        p = footer.paragraphs[0]; p.add_run('Página '); append_field(p, 'PAGE', '1')
        p.add_run(' de '); append_field(p, 'NUMPAGES', '3')
    for page in range(3):
        if page: doc.add_page_break()
        doc.add_heading('Controle sintético '+str(page+1), 1)
        doc.add_paragraph('MARCADOR_EDITAVEL: conteúdo fictício para controle de campo.')
        table = doc.add_table(rows=3, cols=2); table.style = 'Table Grid'
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells): cell.text = f'Célula sintética {page}-{i}-{j}'
        doc.add_paragraph('Parágrafo adicional de comprimento variado. '*10)
    target.parent.mkdir(parents=True, exist_ok=True); doc.save(target)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--complex-savedate', action='store_true'); args = parser.parse_args()
    make(args.out, args.complex_savedate)
