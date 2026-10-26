import sys
import json
import tempfile
import unittest
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from ooxml_fidelity import analyze,UnsafeInput,load_package,safe_xml,W
from fidelity_report import synthesize,pdf_differences
import fitz

class EditorialTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)/'.private'
        self.root.mkdir()
        doc=Document()
        doc.add_paragraph('Texto fictício antes')
        doc.add_paragraph('O segundo permanece.')
        table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='Coluna um'
        table.cell(0,1).text='Coluna dois'
        doc.sections[0].header.paragraphs[0].add_run('Cabeçalho da fixture')
        doc.save(self.root/'source.docx')
        self.source=self.root/'source.docx'
        self.expect={'paragraphIndex':0,'oldText':'Texto fictício antes','newText':'Texto fictício depois'}
        self.candidate=self.root/'candidate.docx'
        self.edit_text()
    def repack(self, mutator):
        with ZipFile(self.source) as z:
            data={n:z.read(n) for n in z.namelist()}
        mutator(data)
        with ZipFile(self.candidate,'w',ZIP_DEFLATED) as z:
            for name,value in data.items():z.writestr(name,value)
    def edit_text(self):
        def mutate(data):
            x=data['word/document.xml']
            self.assertIn(b'Texto fict',x)
            data['word/document.xml']=x.replace(b'Texto fict\xc3\xadcio antes',b'Texto fict\xc3\xadcio depois')
        self.repack(mutate)
    def test_text_edit_prequalified(self):
        r=analyze(self.source,self.candidate,self.expect)
        self.assertEqual(r['outcome'],'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL')
        self.assertFalse(r['mayPromote'])
        self.assertTrue(r['expectedTextEdit']['exact'])
    def test_missing_expected_change_holds(self):
        r=analyze(self.source,self.candidate)
        self.assertEqual(r['outcome'],'QUARANTINE_REQUIRED')
    def test_wrong_expectation_holds(self):
        p=dict(self.expect,newText='Outra coisa')
        self.assertIn('UNEXPECTED_TEXT_CHANGE',analyze(self.source,self.candidate,p)['blockers'])
    def test_added_footnote_part_holds(self):
        def mutate(d):d['word/footnotes.xml']=b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('PART_SET_CHANGED',r['blockers'])
        self.assertIn('word/footnotes.xml',r['addedParts'])
    def test_removed_styles_holds(self):
        def mutate(d):d.pop('word/styles.xml')
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('PART_SET_CHANGED',r['blockers'])
    def test_font_change_holds(self):
        def mutate(d):
            d['word/styles.xml']=d['word/styles.xml'].replace(b'Normal',b'Normax',1)
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('NON_MAIN_PART_CHANGED',r['blockers'])
        self.assertTrue(r['signals']['stylesOrSettingsModified'])
    def test_table_geometry_holds(self):
        def mutate(d):
            d['word/document.xml']=d['word/document.xml'].replace(b'w:type="dxa"',b'w:type="pct"',1)
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('MAIN_STRUCTURE_CHANGED',r['blockers'])
    def test_header_edit_holds(self):
        def mutate(d):
            name=next(n for n in d if n.startswith('word/header') and n.endswith('.xml'))
            d[name]=d[name].replace('Cabeçalho'.encode(), 'Cabecalho'.encode())
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('NON_MAIN_PART_CHANGED',r['blockers'])
    def test_field_rewrite_holds(self):
        def mutate(d):
            xml=d['word/document.xml']
            self.assertIn(b'</w:body>',xml)
            xml=xml.replace(b'</w:body>',b'<w:p><w:fldSimple w:instr="SAVEDATE"/></w:p></w:body>')
            d['word/document.xml']=xml
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertFalse(r['signals']['fieldInstructionsEqual'])
    def test_settings_modified_holds(self):
        def mutate(d):d['word/settings.xml']=d['word/settings.xml'].replace(b'</w:settings>',b'<w:updateFields w:val="true"/></w:settings>')
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('NON_MAIN_PART_CHANGED',r['blockers'])
    def test_media_byte_change_holds(self):
        def mutate(d):d['word/media/image1.png']=b'foo'
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertTrue(r['signals']['mediaChanged'])
    def test_xml_dtd_rejected(self):
        def mutate(d):d['word/document.xml']=b'<!DOCTYPE word [ <!ENTITY hack SYSTEM "file:///etc/passwd"> ]><word>&hack;</word>'
        self.repack(mutate)
        with self.assertRaises(UnsafeInput):analyze(self.source,self.candidate,self.expect)
    def test_traversal_zip_rejected(self):
        def mutate(d):d['../escape.xml']=b'<x/>'
        self.repack(mutate)
        with self.assertRaises(UnsafeInput):analyze(self.source,self.candidate,self.expect)
    def test_nonprivate_path_rejected(self):
        with self.assertRaises(UnsafeInput):load_package(Path(self.tmp.name)/'foo.docx')
    def test_multiple_unexpected_changes_holds(self):
        def mutate(d):d['word/document.xml']=d['word/document.xml'].replace(b'O segundo permanece.',b'O segundo mudou.')
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertIn('UNEXPECTED_TEXT_CHANGE',r['blockers'])
    def test_only_metadata_and_edit_prequalifies(self):
        def mutate(d):
            d['word/document.xml']=d['word/document.xml'].replace('antes'.encode(), 'depois'.encode(),1)
            d['docProps/core.xml']=d['docProps/core.xml'].replace(b'</cp:coreProperties>',b'<cp:lastModifiedBy>Fixture</cp:lastModifiedBy></cp:coreProperties>')
        self.repack(mutate)
        r=analyze(self.source,self.candidate,self.expect)
        self.assertEqual(r['outcome'],'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL')
    def test_fragment_edit_prequalified(self):
        expect={'paragraphIndex':0,'oldFragment':'antes','newFragment':'depois'}
        r=analyze(self.source,self.candidate,expect)
        self.assertTrue(r['expectedTextEdit']['exact'])
        self.assertFalse(r['mayPromote'])
    def test_fragment_ambiguity_holds(self):
        expect={'paragraphIndex':0,'oldFragment':'o','newFragment':'a'}
        r=analyze(self.source,self.candidate,expect)
        self.assertIn('UNEXPECTED_TEXT_CHANGE',r['blockers'])
    def test_invalid_expected_key_rejected(self):
        with self.assertRaises(UnsafeInput):analyze(self.source,self.candidate,{'foo':'bar'})
    def test_no_publication_even_prequalified(self):
        r=analyze(self.source,self.candidate,self.expect)
        self.assertFalse(r['mayPromote'])
    def test_word_metrics_nominal(self):
        def para(text='abc',x=72):return dict(text=text,font='Calibri',fontSize=12,after=8,line=13.9,lineRule=0,page=1,x=x,y=80)
        source={'id':'source','wordBuild':'123','wordVersion':'16','pages':1,'fields':[], 'paragraphs':[para('antes'),para()]}
        candidate={'id':'candidate','wordBuild':'123','wordVersion':'16','pages':1,'fields':[], 'paragraphs':[para('depois'),para()]}
        pdf={'sourcePages':1,'candidatePages':1,'paper':[{'referencePt':[612,792],'candidatePt':[612,792]}], 'rawPixelDifferences':[]}
        oo={'outcome':'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL','blockers':[]}
        result=synthesize(oo,source,candidate,pdf,changed_word_index=0)
        self.assertEqual(result['decision'],'SYNTHETIC_WORD_LAYOUT_PREQUALIFIED_MANUAL_VISUAL_REVIEW_REQUIRED')
        self.assertFalse(result['canPromote'])
    def test_word_auto_detect_unique_change(self):
        def para(text):return dict(text=text,font='Calibri',fontSize=12,after=8,line=13.9,lineRule=0,page=1,x=72,y=80)
        reference={'id':'source','wordBuild':'123','wordVersion':'16','pages':1,'fields':[], 'paragraphs':[para('old'),para('same')]}
        candidate={'id':'candidate','wordBuild':'123','wordVersion':'16','pages':1,'fields':[], 'paragraphs':[para('new'),para('same')]}
        pdf={'sourcePages':1,'candidatePages':1,'paper':[{'referencePt':[612,792],'candidatePt':[612,792]}], 'rawPixelDifferences':[]}
        oo={'outcome':'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL','blockers':[]}
        self.assertTrue(synthesize(oo,reference,candidate,pdf)['metrics']['wordEditPositionVerified'])
    def test_word_metrics_reject_drift(self):
        base=dict(text='same',font='Calibri',fontSize=12,after=8,line=13.9,lineRule=0,page=1,x=72,y=80)
        r={'id':'source','wordBuild':'123','wordVersion':'16','pages':1,'fields':[],'paragraphs':[dict(base,text='old'),dict(base)]}
        c={'id':'candidate','wordBuild':'123','wordVersion':'16','pages':1,'fields':[],'paragraphs':[dict(base,text='new'),dict(base,x=77)]}
        pdf={'sourcePages':1,'candidatePages':1,'paper':[{'referencePt':[612,792],'candidatePt':[612,792]}],'rawPixelDifferences':[]}
        oo={'outcome':'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL','blockers':[]}
        result=synthesize(oo,r,c,pdf,changed_word_index=0)
        self.assertEqual(result['decision'],'HOLD_QUARANTINE')
        self.assertIn('unchangedAnchorsWithinTolerance',result['blockers'])
    def test_word_metrics_reject_missing_change(self):
        src={'id':'source','wordBuild':'123','wordVersion':'16','pages':1,'fields':[], 'paragraphs':[{'text':'no change','x':0,'y':0}]}
        cand={'id':'candidate','wordBuild':'123','wordVersion':'16','pages':1,'fields':[], 'paragraphs':[{'text':'no change','x':0,'y':0}]}
        pdf={'sourcePages':1,'candidatePages':1,'paper':[],'rawPixelDifferences':[]}
        oo={'outcome':'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL','blockers':[]}
        self.assertIn('wordEditPositionVerified',synthesize(oo,src,cand,pdf,changed_word_index=0)['blockers'])
    def test_word_layout_changed_during_pdf_holds(self):
        def para(text):return dict(text=text,font='Cambria',fontSize=11,after=8,line=13,lineRule=0,page=1,x=72,y=80)
        src={'id':'source','wordBuild':'123','wordVersion':'16','pages':1,'fields':[],
             'paragraphs':[para('old'),para('same')],'layoutChangedDuringPdf':False}
        cand=dict(src,id='candidate',paragraphs=[para('new'),para('same')],layoutChangedDuringPdf=True)
        pdf={'sourcePages':1,'candidatePages':1,'paper':[{'referencePt':[612,792],'candidatePt':[612,792]}], 'rawPixelDifferences':[]}
        oo={'outcome':'STRUCTURAL_TEXT_ONLY_PREQUALIFIED_NOT_FIDELITY_APPROVAL','blockers':[]}
        result=synthesize(oo,src,cand,pdf)
        self.assertEqual(result['decision'],'HOLD_QUARANTINE')
        self.assertIn('wordPdfLayoutStateStable',result['blockers'])
    def test_pdf_metrics(self):
        a=self.root/'a.pdf';b=self.root/'b.pdf'
        for p,text in ((a,'A'),(b,'B')):
            d=fitz.open();pg=d.new_page();pg.insert_text((72,72),text);d.save(p);d.close()
        metrics=pdf_differences(a,b)
        self.assertEqual(metrics['sourcePages'],1)
        self.assertGreater(metrics['rawPixelDifferences'][0]['changedPixels'],0)

if __name__=='__main__':unittest.main()
