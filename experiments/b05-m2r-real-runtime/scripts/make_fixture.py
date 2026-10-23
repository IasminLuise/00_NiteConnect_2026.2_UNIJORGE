#!/usr/bin/env python3
"""Synthetic B05-M2R document. Does NOT read official models or user uploads."""
from argparse import ArgumentParser
from pathlib import Path
from docx import Document
from docx.shared import Pt


def make(target: Path) -> None:
    path = target.resolve()
    if '.private' not in path.parts or path.suffix.lower() != '.docx':
        raise ValueError('Output must be a new DOCX under .private')
    if path.exists():
        raise FileExistsError('Never overwrite a fixture')
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading('B05 M2R — Documento fictício', level=1)
    for i in range(3):
        p = doc.add_paragraph(f'Registro de teste {i + 1}: somente conteúdo inventado.')
        p.paragraph_format.space_after = Pt(8)
    doc.add_paragraph('EDITE AQUI: texto inventado para verificar salvamento real.')
    tbl = doc.add_table(rows=2, cols=2)
    for i, row in enumerate(tbl.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f'Dado sintético {i}-{j}'
    doc.save(path)
    print('Fixture sintética criada em pasta privada (não publicar).')


if __name__ == '__main__':
    parser = ArgumentParser()
    parser.add_argument('--out', type=Path, required=True)
    options = parser.parse_args()
    make(options.out)
